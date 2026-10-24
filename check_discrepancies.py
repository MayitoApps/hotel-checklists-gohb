import openpyxl, os, sys

sys.stdout.reconfigure(encoding='utf-8')

files = {
    'Ajusco': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\_Checklist_Actividades_Puestos_OHB_Ajusco.xlsx',
    'Valle de Bravo': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_OHB_Valle.xlsx',
    'Tepoztlan': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_OHB_Tepoztlan.xlsx',
    'Centro Cultural Quinta Soledad': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_CCQS.xlsx',
    'Frontera 49 Roma Norte': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_F49RN.xlsx'
}

seen = set()

for hotel, path in files.items():
    wb = openpyxl.load_workbook(path, data_only=True)
    for sname in wb.sheetnames:
        sheet = wb[sname]
        current_banner = ""
        for r in range(1, sheet.max_row + 1):
            c1 = str(sheet.cell(r, 1).value or '').strip()
            c2 = str(sheet.cell(r, 2).value or '').strip()
            c3 = str(sheet.cell(r, 3).value or '').strip()
            c4 = str(sheet.cell(r, 4).value or '').strip()
            
            if r == 1:
                continue
            if c1 and not c2 and not c3 and not c4:
                if 'estatus' not in c1.lower() and 'checklist' not in c1.lower():
                    current_banner = c1
            elif c4 and 'actividad' not in c4.lower():
                key = (hotel, sname, current_banner, c2)
                if key not in seen:
                    seen.add(key)
                    print(f"[{hotel}] Sheet: {sname:20} | Banner: {current_banner:25} | Col B (Puesto): {c2}")
