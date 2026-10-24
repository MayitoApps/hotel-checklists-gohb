import openpyxl, os, json

files = {
    'Ajusco': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\_Checklist_Actividades_Puestos_OHB_Ajusco.xlsx',
    'Valle de Bravo': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_OHB_Valle.xlsx',
    'Tepoztlan': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_OHB_Tepoztlan.xlsx',
    'Centro Cultural Quinta Soledad': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_CCQS.xlsx',
    'Frontera 49 Roma Norte': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_F49RN.xlsx'
}

analysis = {}

for hotel, path in files.items():
    wb = openpyxl.load_workbook(path, data_only=True)
    analysis[hotel] = {}
    for sname in wb.sheetnames:
        sheet = wb[sname]
        rows = list(sheet.iter_rows(values_only=True))
        
        # let's analyze row by row
        header_row_idx = None
        for idx, r in enumerate(rows):
            line_str = " ".join([str(c) for c in r if c is not None]).lower()
            if 'actividad' in line_str or 'tarea' in line_str:
                header_row_idx = idx
                break
        
        tasks = []
        puestos_in_sheet = set()
        categories_in_sheet = set()
        
        if header_row_idx is not None:
            # find which column is what
            header = [str(c).strip() if c is not None else '' for c in rows[header_row_idx]]
            
            # Identify columns
            status_col = 0
            puesto_col = 1
            cat_col = 2
            task_col = 3
            
            for c_i, h in enumerate(header):
                h_l = h.lower()
                if 'puesto' in h_l:
                    puesto_col = c_i
                elif 'categor' in h_l or 'área' in h_l or 'area' in h_l:
                    cat_col = c_i
                elif 'actividad' in h_l or 'tarea' in h_l:
                    task_col = c_i
            
            for r in rows[header_row_idx + 1:]:
                # check if there is a task
                if len(r) > task_col and r[task_col] is not None and str(r[task_col]).strip():
                    task_desc = str(r[task_col]).strip()
                    puesto_val = str(r[puesto_col]).strip() if len(r) > puesto_col and r[puesto_col] is not None else ''
                    cat_val = str(r[cat_col]).strip() if len(r) > cat_col and r[cat_col] is not None else ''
                    
                    if puesto_val:
                        puestos_in_sheet.add(puesto_val)
                    if cat_val:
                        categories_in_sheet.add(cat_val)
                    
                    tasks.append({
                        'puesto': puesto_val,
                        'categoria': cat_val,
                        'tarea': task_desc,
                        'raw_row': [str(x) if x is not None else '' for x in r if x is not None]
                    })
        
        # Check title rows above header
        pre_headers = []
        if header_row_idx is not None:
            for r in rows[:header_row_idx]:
                vals = [str(x).strip() for x in r if x is not None and str(x).strip()]
                if vals:
                    pre_headers.append(" | ".join(vals))
                    
        analysis[hotel][sname] = {
            'header_row_index': header_row_idx,
            'pre_headers': pre_headers,
            'puestos_found': sorted(list(puestos_in_sheet)),
            'categories_found': sorted(list(categories_in_sheet)),
            'total_tasks': len(tasks),
            'tasks': tasks
        }

with open('full_analysis.json', 'w', encoding='utf-8') as f:
    json.dump(analysis, f, ensure_ascii=False, indent=2)

print("Full analysis saved to full_analysis.json")
