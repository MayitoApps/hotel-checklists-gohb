import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\_Checklist_Actividades_Puestos_OHB_Ajusco.xlsx', data_only=True)
sheet = wb['MANTENIMIENTO']
print("=== AJUSCO MANTENIMIENTO ===")
for r in range(1, sheet.max_row + 1):
    vals = [sheet.cell(r, c).value for c in range(1, 5)]
    if any(vals):
        print(f"R{r}: {vals}")

sheet = wb['Origen Centro Holistico']
print("\n=== AJUSCO ORIGEN CENTRO HOLISTICO ===")
for r in range(1, sheet.max_row + 1):
    vals = [sheet.cell(r, c).value for c in range(1, 5)]
    if any(vals):
        print(f"R{r}: {vals}")
