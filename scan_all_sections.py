import openpyxl, os

files = {
    'Ajusco': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\_Checklist_Actividades_Puestos_OHB_Ajusco.xlsx',
    'Valle de Bravo': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_OHB_Valle.xlsx',
    'Tepoztlan': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_OHB_Tepoztlan.xlsx',
    'Centro Cultural Quinta Soledad': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_CCQS.xlsx',
    'Frontera 49 Roma Norte': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_F49RN.xlsx'
}

for hotel, path in files.items():
    print(f"\n=======================================================")
    print(f"HOTEL: {hotel}")
    print(f"=======================================================")
    wb = openpyxl.load_workbook(path, data_only=True)
    for sname in wb.sheetnames:
        sheet = wb[sname]
        print(f"\n--- SHEET: {sname} (Rows: {sheet.max_row}) ---")
        for r in range(1, sheet.max_row + 1):
            row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
            non_empty = [(c_idx + 1, v) for c_idx, v in enumerate(row_vals) if v is not None and str(v).strip()]
            if non_empty:
                # check if it's a section header or regular row
                val_str = " | ".join([f"C{col}:{val}" for col, val in non_empty])
                if len(non_empty) == 1 and non_empty[0][0] == 1:
                    print(f"  [SECTION / TITLE] Row {r}: {non_empty[0][1]}")
                elif any('estatus' in str(v).lower() for _, v in non_empty):
                    print(f"  [HEADER] Row {r}: {val_str}")
                else:
                    # sample task row
                    puesto = sheet.cell(r, 2).value
                    cat = sheet.cell(r, 3).value
                    task = sheet.cell(r, 4).value
                    # print row number and brief summary
                    print(f"    Row {r}: Puesto='{puesto}' | Cat='{cat}' | Tarea='{str(task)[:45]}...'")
