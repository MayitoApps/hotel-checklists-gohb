import openpyxl, json, os

files = {
    'Orgánico Hotel Boutique Ajusco': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\_Checklist_Actividades_Puestos_OHB_Ajusco.xlsx',
    'Orgánico Hotel Boutique Valle de Bravo': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_OHB_Valle.xlsx',
    'Orgánico Hotel Boutique Tepoztlán': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_OHB_Tepoztlan.xlsx',
    'Centro Cultural Quinta Soledad': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_CCQS.xlsx',
    'Frontera 49 Roma Norte': r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\Checklist_Actividades_Puestos_F49RN.xlsx'
}

catalog = {}

for hotel_name, file_path in files.items():
    wb = openpyxl.load_workbook(file_path, data_only=True)
    catalog[hotel_name] = {}
    
    for sname in wb.sheetnames:
        sheet = wb[sname]
        
        # We will iterate row by row to detect sections
        current_section = sname
        
        for r in range(1, sheet.max_row + 1):
            c1 = sheet.cell(r, 1).value
            c2 = sheet.cell(r, 2).value
            c3 = sheet.cell(r, 3).value
            c4 = sheet.cell(r, 4).value
            
            c1_s = str(c1 or '').strip()
            c2_s = str(c2 or '').strip()
            c3_s = str(c3 or '').strip()
            c4_s = str(c4 or '').strip()
            
            # Skip title row 1
            if r == 1:
                continue
                
            # Check if this row is a section title (e.g., 'Mesero A', 'Chef', 'Jefe de mantenimiento')
            if c1_s and not c2_s and not c3_s and not c4_s:
                if 'estatus' not in c1_s.lower() and 'checklist' not in c1_s.lower():
                    current_section = c1_s
                    continue
                    
            # Skip table header row
            if 'estatus' in c1_s.lower() or 'actividad / tarea' in c4_s.lower():
                continue
                
            # If c4 is present and has task text
            if c4_s and c4_s.lower() != 'none':
                # Determine the position name:
                # If section banner is specific (e.g., 'TERAPEUTAS INTERNAS', 'Jefe de mantenimiento', 'Mesero A', 'Ayudantes de cocina', 'Panadero', 'Auxiliar de panadería')
                # Let's record both the section banner and column 2 value
                
                # Priority: if current_section has a clear specific role name, we use it, but let's keep track
                puesto_final = current_section
                # if current_section was just sheet name like 'Gerente' and c2 has 'Gerencia General', we can refine
                if current_section in [sname, ''] and c2_s:
                    puesto_final = c2_s
                elif current_section.lower() in ['recepcionista', 'recepcionista a'] and c2_s:
                    puesto_final = c2_s
                    
                # Normalize spaces
                puesto_final = " ".join(puesto_final.split())
                
                if puesto_final not in catalog[hotel_name]:
                    catalog[hotel_name][puesto_final] = {
                        'area': sname,
                        'puesto_col_b': c2_s,
                        'tasks': []
                    }
                    
                # Check if task already exists in this puesto (to avoid duplicates if the section was pasted twice)
                existing_tasks = [t['tarea'] for t in catalog[hotel_name][puesto_final]['tasks']]
                if c4_s not in existing_tasks:
                    catalog[hotel_name][puesto_final]['tasks'].append({
                        'categoria': c3_s if c3_s else 'General',
                        'tarea': c4_s
                    })

with open('catalog_positions.json', 'w', encoding='utf-8') as f:
    json.dump(catalog, f, ensure_ascii=False, indent=2)

print("Exported catalog_positions.json successfully")
