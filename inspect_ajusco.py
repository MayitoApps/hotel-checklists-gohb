import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\maval\OneDrive\Documentos\Check-lsit app\_Checklist_Actividades_Puestos_OHB_Ajusco.xlsx', data_only=True)
for sname in wb.sheetnames:
    sheet = wb[sname]
    print(f"=== SHEET: {sname} ===")
    for i in range(1, 10):
        row_vals = [sheet.cell(i, j).value for j in range(1, 8)]
        if any(row_vals):
            print(f"Row {i}: {row_vals}")
    print()
